from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from dotenv import load_dotenv
from google import genai
import os
import time
import json

from pypdf import PdfReader
from pptx import Presentation


# ==========================================================
# MEMORA
# AI-POWERED PERSONALIZED LEARNING & MEMORY PLATFORM
# ==========================================================

load_dotenv()

app = Flask(__name__)
CORS(app)

api_key = os.getenv("GEMINI_API_KEY")

if not api_key:
    print("WARNING: GEMINI_API_KEY is not set.")

client = genai.Client(api_key=api_key)


# ==========================================================
# HELPER - GEMINI AI
# ==========================================================

def generate_ai(prompt):

    for attempt in range(3):

        try:

            response = client.models.generate_content(
                model="gemini-3.6-flash",
                contents=prompt
            )

            return response.text

        except Exception as error:

            error_text = str(error)

            if (
                "429" in error_text
                or "RESOURCE_EXHAUSTED" in error_text
            ):

                print("ERROR: Gemini quota exceeded.")

                raise Exception(
                    "MEMORA AI limit has been reached. Please try again later."
                )

            if "503" in error_text and attempt < 2:

                print("Gemini is busy. Retrying...")

                time.sleep(3)

                continue

            raise error


# ==========================================================
# MAIN AI GENERATOR
# ==========================================================

@app.route("/api/generate", methods=["POST"])
def generate():

    try:

        data = request.get_json(silent=True) or {}

        mode = data.get(
            "mode",
            "explain"
        )


        # ==================================================
        # STUDY PLANNER
        # ==================================================

        if mode == "planner":

            exam_name = str(
                data.get("examName", "")
            ).strip()

            exam_date = str(
                data.get("examDate", "")
            ).strip()

            subjects = str(
                data.get("subjects", "")
            ).strip()

            study_hours = str(
                data.get("studyHours", "")
            ).strip()


            # ------------------------------
            # VALIDATION
            # ------------------------------

            if not subjects:

                return jsonify({
                    "error": "Please enter at least one subject."
                }), 400


            if not exam_name:

                exam_name = "Upcoming Exam"


            if not study_hours:

                study_hours = "2"


            if not exam_date:

                exam_date = "Not specified"


            # ------------------------------
            # CLEAN SUBJECT LIST
            # ------------------------------

            subject_lines = []

            for line in subjects.splitlines():

                line = line.strip()

                if line:

                    subject_lines.append(line)


            if not subject_lines:

                return jsonify({
                    "error": "Please enter at least one subject."
                }), 400


            subject_text = "\n".join(
                f"- {subject}"
                for subject in subject_lines
            )


            # ------------------------------
            # AI PLANNER PROMPT
            # ------------------------------

            planner_prompt = f"""
You are MEMORA, an intelligent college exam study planner.

Create a SPECIFIC, PRACTICAL and SUBJECT-WISE study plan.

EXAM:
{exam_name}

EXAM DATE:
{exam_date}

SUBJECTS:
{subject_text}

DAILY STUDY TIME:
{study_hours} hours


IMPORTANT:

The subject names above may be broad subjects such as:

- Data Structures
- History
- DBMS
- Java
- Artificial Intelligence
- Mathematics

DO NOT treat the broad subject name itself as the study topic.

You must break every subject into smaller academic topics.

For example:

Data Structures can contain topics such as:
- Arrays
- Linked Lists
- Stacks
- Queues
- Trees
- Graphs
- Searching
- Sorting
- Algorithms
- Complexity

History can contain topics such as:
- Historical events
- Causes
- Important dates
- Important people
- Major developments
- Effects
- Comparisons
- Previous-question revision

However, ONLY use topics that are reasonably connected to the subject.
Do not randomly combine unrelated subjects.


CORE RULES:

1. Use ONLY the subjects provided by the student.

2. NEVER make the entire day simply:
   "Study Data Structures"
   or
   "Study History".

3. Every session must contain a SPECIFIC topic.

4. Every session must contain a SPECIFIC activity.

5. Break broad subjects into multiple smaller topics.

6. Do not repeat the same topic on different days unless
   the session is clearly revision, practice, or active recall.

7. Distribute the subjects fairly.

8. Alternate subjects when possible so the student does
   not study the same subject continuously.

9. Start with fundamentals.

10. Move from fundamentals to important concepts,
    applications, examples and problem solving.

11. Include coding/problem-solving practice for technical
    subjects when appropriate.

12. Include dates, events, causes, effects and active recall
    for History when appropriate.

13. Include revision and active recall near the end.

14. Include a mock test or previous-question practice
    near the exam.

15. The plan must fit within the student's daily study time.

16. Each day must have exactly 2 study sessions.

17. Each session should represent roughly half of the
    available daily study time.

18. Use concise topic names.

19. The student should immediately understand WHAT to study
    and WHAT to do.

20. Do not create a generic timetable.

21. Do not invent unrelated subjects.

22. Return ONLY valid JSON.

23. Do NOT use Markdown code fences.


RETURN EXACTLY THIS STRUCTURE:

{{
    "days": [
        {{
            "day": 1,
            "sessions": [
                {{
                    "subject": "Data Structures",
                    "topic": "specific topic",
                    "activity": "specific study activity"
                }},
                {{
                    "subject": "History",
                    "topic": "specific topic",
                    "activity": "specific study activity"
                }}
            ],
            "recall": "short active recall task"
        }}
    ]
}}


IMPORTANT:

The "subject" field must contain the actual subject.

The "topic" field must contain a smaller topic INSIDE that subject.

The "activity" field must explain what the student should actually do.


BAD:

{{
    "subject": "Data Structures",
    "topic": "Data Structures",
    "activity": "Study Data Structures"
}}


GOOD:

{{
    "subject": "Data Structures",
    "topic": "Linked List Insertion and Deletion",
    "activity": "Trace insertion and deletion algorithms and solve 2 practice problems"
}}


Create a logical progression across the days.

If there are many days, use later days for:

- revision
- active recall
- practice questions
- mock tests
- weak-topic revision

Make the plan useful for actual exam preparation.
"""


            try:

                ai_text = generate_ai(
                    planner_prompt
                )

            except Exception as error:

                if (
                    "AI limit" in str(error)
                    or "quota" in str(error).lower()
                ):

                    return jsonify({
                        "error": str(error)
                    }), 429

                print(
                    "PLANNER AI ERROR:",
                    error
                )

                return jsonify({
                    "error":
                    "MEMORA could not create the study plan. Please try again."
                }), 500


            # ------------------------------
            # CLEAN AI JSON
            # ------------------------------

            ai_text = ai_text.strip()


            if ai_text.startswith("```json"):

                ai_text = ai_text[7:]


            elif ai_text.startswith("```"):

                ai_text = ai_text[3:]


            if ai_text.endswith("```"):

                ai_text = ai_text[:-3]


            ai_text = ai_text.strip()


            # ------------------------------
            # PARSE JSON
            # ------------------------------

            try:

                planner_data = json.loads(
                    ai_text
                )

            except json.JSONDecodeError as error:

                print(
                    "PLANNER JSON ERROR:",
                    error
                )

                print(
                    "AI RESPONSE:",
                    ai_text
                )

                return jsonify({
                    "error":
                    "MEMORA received an invalid study plan. Please try again."
                }), 500


            # ------------------------------
            # VALIDATE JSON
            # ------------------------------

            if not isinstance(
                planner_data,
                dict
            ):

                return jsonify({
                    "error":
                    "MEMORA could not create a valid study plan."
                }), 500


            if "days" not in planner_data:

                return jsonify({
                    "error":
                    "MEMORA could not create a valid study plan."
                }), 500


            if not isinstance(
                planner_data["days"],
                list
            ):

                return jsonify({
                    "error":
                    "MEMORA could not create a valid study plan."
                }), 500


            # ------------------------------
            # RETURN PLANNER
            # ------------------------------

            return jsonify({
                "result": planner_data
            })


        # ==================================================
        # NORMAL AI MODES
        # ==================================================

        user_input = str(
            data.get("input", "")
        ).strip()


        if not user_input:

            return jsonify({
                "error": "Please enter something to study."
            }), 400


        # ==================================================
        # EXPLAIN
        # ==================================================

        if mode == "explain":

            prompt = f"""
You are MEMORA, an AI study assistant.

Explain the following topic clearly for a college student.

Use this structure:

## 🧠 Explanation

Give a clear explanation.

## 📌 Important Points

List the important ideas.

## 💡 Example

Give a simple example.

## 🧠 Memory Tip

Give one useful memory trick.

Use simple language.

Do not use LaTeX.

Do not use dollar signs for mathematics.

Topic:
{user_input}
"""


        # ==================================================
        # SUMMARIZE
        # ==================================================

        elif mode == "summarize":

            prompt = f"""
You are MEMORA, an AI study assistant.

Summarize the following study material.

Use this structure:

## 📝 Summary

Give a concise but useful summary.

## ⭐ Key Points

List the most important points.

## 🧠 Remember

Give the most important thing the student should remember.

Do not use LaTeX.

Do not use dollar signs for mathematics.

Study material:
{user_input}
"""


        # ==================================================
        # IMPORTANT POINTS
        # ==================================================

        elif mode == "important":

            prompt = f"""
You are MEMORA, an AI study assistant.

Find the most important information from the study material.

Use this structure:

## ⭐ Important Points

Give the important points.

## 📌 Must Remember

Give the facts or concepts most likely to matter in exams.

## 🧠 Memory Tip

Give one useful memory trick.

Do not use LaTeX.

Do not use dollar signs for mathematics.

Study material:
{user_input}
"""


        # ==================================================
        # QUIZ
        # ==================================================

        elif mode == "quiz":

            prompt = f"""
You are MEMORA, an AI study assistant.

Create 5 useful quiz questions from the following
study material.

Give:

- Question
- 4 options
- Correct answer
- Short explanation

Make the questions useful for exam preparation.

Do not use LaTeX.

Do not use dollar signs for mathematics.

Study material:
{user_input}
"""


        # ==================================================
        # IMPROVE ANSWER
        # ==================================================

        elif mode == "improve":

            prompt = f"""
You are MEMORA, an AI study assistant.

Analyze the student's answer below.

Give:

1. What is correct
2. What needs improvement
3. A better version of the answer
4. One tip for the student

Be supportive and clear.

Do not use LaTeX.

Do not use dollar signs for mathematics.

Student answer:
{user_input}
"""


        # ==================================================
        # I DON'T GET IT
        # ==================================================

        elif mode == "simple":

            prompt = f"""
You are MEMORA's "I DON'T GET IT" mode.

The student is confused and needs the topic explained
as if they are learning it for the FIRST TIME.

IMPORTANT:

- Do NOT give a normal textbook explanation.
- Use very simple everyday words.
- Keep sentences short.
- Avoid unnecessary technical terminology.
- If a technical word is necessary, explain it immediately.
- Make the explanation friendly and encouraging.

Use exactly this structure:

## 😵 Let's Make It Super Easy

### 🧠 In One Sentence

Explain the topic in ONE very simple sentence.

### 🏠 Imagine This

Give a funny or relatable everyday-life analogy.

### 📦 Step-by-Step

Explain how it works in 3–5 very simple steps.

### 🎯 Easy Example

Give one small example that a college student can easily understand.

### 🧠 Remember It Like This

Give one short memory trick.

### ⚡ In Short

Give a final explanation in 1–2 simple sentences.

Do not use LaTeX.

Do not use dollar signs for mathematics.

Topic:
{user_input}
"""


        # ==================================================
        # TEACH IT BACK
        # ==================================================

        elif mode == "teach":

            prompt = f"""
You are MEMORA's "Teach It Back" learning evaluator.

A student is trying to explain a topic in their own words.

Your job is to check how well they understand it.

Be supportive and encouraging.

Do not shame the student for mistakes.

Use exactly this structure:

## 🗣️ Teach It Back Feedback

### 🟢 What You Got Right

Mention the important ideas the student explained correctly.

### 🟡 What Is Missing

Mention important concepts or details that the student did not include.

### 🔴 What Needs Correction

Identify any incorrect statements and explain them simply.

### ✨ Better Explanation

Give a clear and simple version of the topic.

### 🧠 Memory Tip

Give one short trick to remember the important idea.

### 📊 Understanding Level

Give a simple rating:

- Excellent
- Good
- Needs More Practice

Do not use LaTeX.

Do not use dollar signs for mathematics.

Topic:
{user_input}
"""


        # ==================================================
        # QUICK STUDY
        # ==================================================

        elif mode == "quick-study":

            prompt = f"""
You are MEMORA, a friendly AI learning assistant.

The student wants to quickly learn this topic:

{user_input}

Actually TEACH the topic.

Do NOT create a timetable.

Use this structure:

## 🧠 What Is It?

Explain the topic simply.

## 📌 Core Concept

Explain the most important idea.

## ⭐ Important Points

Give the key things the student should remember.

## 💡 Simple Example

Give one easy example.

## ❓ Quick Practice

Give one short practice question.

## 🧠 Memory Trick

Give one easy way to remember it.

## 🔁 Final Recap

Give 3 to 5 short revision points.

Use simple student-friendly language.

Do not use LaTeX.

Do not use dollar signs for mathematical notation.
"""


        else:

            return jsonify({
                "error": "Invalid study mode."
            }), 400


        # ==================================================
        # GENERATE NORMAL AI RESPONSE
        # ==================================================

        try:

            result = generate_ai(
                prompt
            )

        except Exception as error:

            error_text = str(error)

            if (
                "AI limit" in error_text
                or "quota" in error_text.lower()
            ):

                return jsonify({
                    "error": error_text
                }), 429

            print(
                "AI ERROR:",
                error
            )

            return jsonify({
                "error":
                "MEMORA could not generate a response. Please try again."
            }), 500


        return jsonify({
            "result": result
        })


    except Exception as error:

        print(
            "ERROR:",
            error
        )

        return jsonify({
            "error":
            "MEMORA could not generate a response. Please try again."
        }), 500


# ==========================================================
# IMAGE STUDY 📷🧠
# ==========================================================

@app.route("/api/study-image", methods=["POST"])
def study_image():

    try:

        if "image" not in request.files:

            return jsonify({
                "error": "Please upload an image."
            }), 400


        image = request.files["image"]


        if image.filename == "":

            return jsonify({
                "error": "Please choose an image."
            }), 400


        image_data = image.read()


        response = client.models.generate_content(

            model="gemini-3.6-flash",

            contents=[

                """
You are MEMORA, an AI study assistant.

Look at the uploaded study image and help the student
understand the material.

Give the answer using this structure:

## 🧠 What This Image Contains

Briefly explain what the image is about.

## 📚 Simple Explanation

Explain the important information in very simple language.

## ⭐ Important Points

Give the most important points the student should remember.

## 🧠 Memory Tip

Give one easy trick to remember the information.

## ❓ Quick Questions

Give 3 simple questions the student can use to test
their understanding.

Be clear, friendly and useful for a college student.

Do not use LaTeX.
Do not use dollar signs for mathematics.
""",

                genai.types.Part.from_bytes(
                    data=image_data,
                    mime_type=image.content_type
                )

            ]
        )


        return jsonify({
            "result": response.text
        })


    except Exception as error:

        print(
            "IMAGE ERROR:",
            error
        )

        if (
            "429" in str(error)
            or "RESOURCE_EXHAUSTED" in str(error)
        ):

            return jsonify({
                "error":
                "MEMORA AI limit has been reached. Please try again later."
            }), 429


        return jsonify({
            "error":
            "MEMORA could not study this image."
        }), 500


# ==========================================================
# STUDY PDF 📄🧠
# ==========================================================

@app.route("/api/study-pdf", methods=["POST"])
def study_pdf():

    try:

        if "file" not in request.files:

            return jsonify({
                "error": "Please upload a PDF."
            }), 400


        pdf_file = request.files["file"]


        if pdf_file.filename == "":

            return jsonify({
                "error": "Please choose a PDF."
            }), 400


        reader = PdfReader(
            pdf_file
        )

        text = ""


        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:

                text += page_text + "\n"


        if not text.strip():

            return jsonify({
                "error":
                "MEMORA could not find readable text in this PDF."
            }), 400


        return jsonify({
            "result": text
        })


    except Exception as error:

        print(
            "PDF ERROR:",
            error
        )

        return jsonify({
            "error":
            "MEMORA could not read this PDF."
        }), 500


# ==========================================================
# STUDY POWERPOINT 📊🧠
# ==========================================================

@app.route("/api/study-ppt", methods=["POST"])
def study_ppt():

    try:

        if "file" not in request.files:

            return jsonify({
                "error":
                "Please upload a PowerPoint file."
            }), 400


        ppt_file = request.files["file"]


        if ppt_file.filename == "":

            return jsonify({
                "error":
                "Please choose a PowerPoint file."
            }), 400


        presentation = Presentation(
            ppt_file
        )

        text = ""


        for slide_number, slide in enumerate(
            presentation.slides,
            start=1
        ):

            text += (
                f"\n--- Slide {slide_number} ---\n"
            )


            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    if shape.text.strip():

                        text += (
                            shape.text
                            + "\n"
                        )


        if not text.strip():

            return jsonify({
                "error":
                "MEMORA could not find readable text in this presentation."
            }), 400


        return jsonify({
            "result": text
        })


    except Exception as error:

        print(
            "PPT ERROR:",
            error
        )

        return jsonify({
            "error":
            "MEMORA could not read this presentation."
        }), 500


# ==========================================================
# STUDY TEXT FILE 📝🧠
# ==========================================================

@app.route("/api/study-txt", methods=["POST"])
def study_txt():

    try:

        if "file" not in request.files:

            return jsonify({
                "error":
                "Please upload a text file."
            }), 400


        txt_file = request.files["file"]


        if txt_file.filename == "":

            return jsonify({
                "error":
                "Please choose a text file."
            }), 400


        text = txt_file.read().decode(
            "utf-8",
            errors="replace"
        )


        if not text.strip():

            return jsonify({
                "error":
                "MEMORA could not find any readable text in this file."
            }), 400


        return jsonify({
            "result": text
        })


    except Exception as error:

        print(
            "TXT ERROR:",
            error
        )

        return jsonify({
            "error":
            "MEMORA could not read this text file."
        }), 500


# ==========================================================
# FRONTEND ROUTES
# ==========================================================

@app.route("/")
def home():

    return send_file(
        "index.html"
    )


@app.route("/style.css")
def css():

    return send_file(
        "style.css"
    )


@app.route("/script.js")
def js():

    return send_file(
        "script.js"
    )


# ==========================================================
# START MEMORA
# ==========================================================

if __name__ == "__main__":

    app.run(
        debug=True
    )